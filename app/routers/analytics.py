"""Analytics API router."""
from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func
import csv
import io
import json
from pathlib import Path
from datetime import datetime

from app.database import get_db
from app.models.interview import Interview
from sqlalchemy import func
from app.models.contact import Contact
from app.dependencies import get_current_user
from app.models.user import User

router = APIRouter(prefix="/api/analytics", tags=["analytics"])


def load_form_schema():
    """Load form schema to get field labels."""
    try:
        schema_path = Path(__file__).parent.parent.parent.parent / "public" / "form_schema_with_nav.json"
        if schema_path.exists():
            with open(schema_path, 'r', encoding='utf-8') as f:
                return json.load(f)
    except Exception:
        pass
    return None


def get_field_map(schema):
    """Build field_id -> {label, options} map from schema."""
    field_map = {}
    if not schema or 'pages' not in schema:
        return field_map
    for page in schema['pages']:
        for field in page.get('fields', []):
            opts = {}
            for o in field.get('options', []):
                opts[o['value']] = o['label']
            field_map[field['id']] = {
                'label': field.get('label', field['id']),
                'options': opts,
                'type': field.get('type', 'text'),
            }
    return field_map

@router.get("/interview-statistics/")
async def interview_statistics(
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Return statistics for interviews."""
    total_interviews_result = await db.execute(select(func.count(Interview.id)))
    total_interviews = total_interviews_result.scalar() or 0

    completed_result = await db.execute(
        select(func.count(Interview.id)).where(Interview.completed_at.isnot(None))
    )
    completed_interviews = completed_result.scalar() or 0

    in_progress_result = await db.execute(
        select(func.count(Interview.id)).where(
            (Interview.started_at.isnot(None)) & (Interview.completed_at.is_(None))
        )
    )
    in_progress = in_progress_result.scalar() or 0

    completion_rate = round((completed_interviews / total_interviews * 100), 1) if total_interviews > 0 else 0

    return {
        "total_interviews": total_interviews,
        "completed_interviews": completed_interviews,
        "in_progress_interviews": in_progress,
        "completion_rate": completion_rate,
    }

def resolve_value(field_info, value):
    """Resolve a value to its human-readable label if it's a select option."""
    if value is None or value == '':
        return ''
    if field_info and field_info.get('options'):
        if isinstance(value, list):
            return '; '.join(field_info['options'].get(str(v), str(v)) for v in value)
        return field_info['options'].get(str(value), str(value))
    return str(value)


@router.get("/summary/")
async def get_analytics_summary(
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Get summary statistics for the analytics dashboard."""
    total_contacts_result = await db.execute(select(func.count(Contact.id)))
    total_contacts = total_contacts_result.scalar() or 0

    total_interviews_result = await db.execute(select(func.count(Interview.id)))
    total_interviews = total_interviews_result.scalar() or 0

    completed_result = await db.execute(
        select(func.count(Interview.id)).where(Interview.completed_at.isnot(None))
    )
    completed_interviews = completed_result.scalar() or 0

    in_progress_result = await db.execute(
        select(func.count(Interview.id)).where(
            (Interview.started_at.isnot(None)) & (Interview.completed_at.is_(None))
        )
    )
    in_progress = in_progress_result.scalar() or 0

    completion_rate = round((completed_interviews / total_interviews * 100), 1) if total_interviews > 0 else 0

    return {
        "total_contacts": total_contacts,
        "total_interviews": total_interviews,
        "completed_interviews": completed_interviews,
        "in_progress_interviews": in_progress,
        "completion_rate": completion_rate,
    }


@router.get("/enhanced-download/")
async def enhanced_download(
    report: str = Query("roster", regex="^(roster|interview_summary|detailed)$"),
    format: str = Query("csv", regex="^(csv|xlsx)$"),
    limit: int = Query(None, ge=1, le=100000),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Download reports as CSV or XLSX."""

    # --- Gather data ---
    contacts_result = await db.execute(select(Contact).order_by(Contact.id))
    contacts = contacts_result.scalars().all()

    interviews_result = await db.execute(
        select(Interview).order_by(Interview.contact_id, Interview.round_number)
    )
    interviews = interviews_result.scalars().all()

    # Build contact map
    contact_map = {}
    for c in contacts:
        contact_map[c.id] = c

    # Build interviews-by-contact map
    interviews_by_contact: dict[int, list] = {}
    for iv in interviews:
        interviews_by_contact.setdefault(iv.contact_id, []).append(iv)

    schema = load_form_schema()
    field_map = get_field_map(schema)

    # Ordered field IDs from schema
    ordered_fields = []
    if schema and 'pages' in schema:
        for page in schema['pages']:
            for field in page.get('fields', []):
                ordered_fields.append(field['id'])

    # --- Build rows ---
    if report == "roster":
        headers = [
            "ID", "Name", "Phone", "Serial Number", "CUID", "Ticket Number",
            "Status", "Location", "Study Arm", "Method",
            "Batch Number", "Batch Serial", "Interview Count", "Created At"
        ]
        rows = []
        for c in contacts:
            rows.append([
                c.id, c.name, c.phone, c.serial_number, c.cuid or '', c.ticket_number or '',
                c.status, c.location or '', c.study_arm or '', c.method or '',
                c.batch_number or '', c.batch_serial_number or '',
                c.interview_count or 0,
                c.created_at.isoformat() if c.created_at else '',
            ])
        filename = "contact-roster"

    elif report == "interview_summary":
        headers = [
            "Contact ID", "Name", "CUID", "Ticket Number", "Status", "Location",
            "Round 1 Status", "Round 1 Completed",
            "Round 2 Status", "Round 2 Completed",
            "Round 3 Status", "Round 3 Completed",
            "Round 4 Status", "Round 4 Completed",
        ]
        rows = []
        for c in contacts:
            ivs = interviews_by_contact.get(c.id, [])
            iv_by_round = {iv.round_number: iv for iv in ivs}
            row = [
                c.id, c.name, c.cuid or '', c.ticket_number or '', c.status, c.location or '',
            ]
            for rn in [1, 2, 3, 4]:
                iv = iv_by_round.get(rn)
                if iv:
                    status = 'completed' if iv.completed_at else ('in_progress' if iv.started_at else 'not_started')
                    completed = iv.completed_at.isoformat() if iv.completed_at else ''
                else:
                    status = 'not_started'
                    completed = ''
                row.extend([status, completed])
            rows.append(row)
        filename = "interview-summary"

    else:  # detailed
        # Build headers: contact fields + each form field with human-readable label
        contact_headers = [
            "Interview ID", "Contact ID", "Name", "CUID", "Ticket Number",
            "Round", "Status", "Started At", "Completed At",
        ]
        # Use short label (first 80 chars) for column headers
        form_headers = []
        for fid in ordered_fields:
            info = field_map.get(fid, {})
            label = info.get('label', fid)
            short = label[:80] + '...' if len(label) > 80 else label
            form_headers.append(short)

        headers = contact_headers + form_headers
        rows = []
        for iv in interviews:
            if iv.completed_at is None:
                continue  # Only completed interviews
            c = contact_map.get(iv.contact_id)
            if not c:
                continue
            row = [
                iv.id, c.id, c.name, c.cuid or '', c.ticket_number or '',
                iv.round_number,
                'completed' if iv.completed_at else 'in_progress',
                iv.started_at.isoformat() if iv.started_at else '',
                iv.completed_at.isoformat() if iv.completed_at else '',
            ]
            xform = iv.xform_data or {}
            for fid in ordered_fields:
                val = xform.get(fid, '')
                row.append(resolve_value(field_map.get(fid), val))
            rows.append(row)
        filename = "interview-details"

    # Apply limit
    if limit and len(rows) > limit:
        rows = rows[:limit]

    # --- Generate output ---
    timestamp = datetime.utcnow().strftime("%Y-%m-%d")

    if format == "xlsx":
        import openpyxl
        from openpyxl.styles import Font, PatternFill

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = report.replace('_', ' ').title()

        # Header row
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill

        # Data rows
        for row_idx, row in enumerate(rows, 2):
            for col_idx, val in enumerate(row, 1):
                ws.cell(row=row_idx, column=col_idx, value=val)

        # Auto-width (cap at 50)
        for col_idx, h in enumerate(headers, 1):
            max_len = len(str(h))
            for row in rows[:50]:  # Sample first 50 rows
                if col_idx <= len(row):
                    max_len = max(max_len, len(str(row[col_idx - 1])))
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = min(max_len + 2, 50)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}-{timestamp}.xlsx"},
        )
    else:
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(headers)
        writer.writerows(rows)

        return StreamingResponse(
            iter([output.getvalue()]),
            media_type="text/csv",
            headers={"Content-Disposition": f"attachment; filename={filename}-{timestamp}.csv"},
        )
